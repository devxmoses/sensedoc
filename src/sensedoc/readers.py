"""Turn uploaded files into text and/or images the LLM can read, plus where each line sits."""

from dataclasses import dataclass, field
from pathlib import Path

import docx
import pandas as pd
import pymupdf

from sensedoc import config
from sensedoc.ocr import Line, ocr_image, pdf_page_lines

MIN_TEXT_CHARS = 50  # a PDF page with less text than this is probably a scan

IMAGE_TYPES = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".webp": "image/webp",
}


@dataclass
class Image:
    data: bytes
    mime_type: str
    page: int = 1


@dataclass
class ParsedDocument:
    text: str = ""  # real text from the file (digital PDF, Excel, Word)
    images: list[Image] = field(default_factory=list)  # pages that are pictures (photos, scans)
    page_count: int = 1
    lines: list[Line] = field(default_factory=list)  # every line of text with its position

    @property
    def ocr_text(self) -> str:
        """Text that OCR read from the page images (may contain reading errors)."""
        return "\n".join(line.text for line in self.lines if line.from_ocr)


def read_pdf(path: Path, ocr: bool) -> ParsedDocument:
    parsed = ParsedDocument()
    texts = []
    with pymupdf.open(path) as pdf:
        parsed.page_count = len(pdf)
        for number, page in enumerate(pdf, start=1):
            page_text = page.get_text().strip()
            if len(page_text) >= MIN_TEXT_CHARS:
                texts.append(f"--- Page {number} ---\n{page_text}")
                parsed.lines += pdf_page_lines(page, number)
            else:
                # no real text: probably a scan, so send a picture of the page
                png = page.get_pixmap(dpi=200).tobytes("png")
                parsed.images.append(Image(png, "image/png", number))
                if ocr:
                    parsed.lines += ocr_image(png, number)
    parsed.text = "\n\n".join(texts)
    return parsed


def read_image(path: Path, ocr: bool) -> ParsedDocument:
    data = path.read_bytes()
    parsed = ParsedDocument(images=[Image(data, IMAGE_TYPES[path.suffix.lower()])])
    if ocr:
        parsed.lines = ocr_image(data, 1)
    return parsed


def read_spreadsheet(path: Path) -> ParsedDocument:
    if path.suffix.lower() == ".csv":
        sheets = {"Sheet1": pd.read_csv(path)}
    else:
        sheets = pd.read_excel(path, sheet_name=None)  # None means "every sheet"
    parts = [f"--- Sheet: {name} ---\n{df.to_csv(index=False)}" for name, df in sheets.items()]
    return ParsedDocument(text="\n\n".join(parts))


def read_docx(path: Path) -> ParsedDocument:
    document = docx.Document(path)
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return ParsedDocument(text="\n".join(parts))


def read_document(path: str | Path, ocr: bool | None = None) -> ParsedDocument:
    """Pick the right reader based on the file extension."""
    path = Path(path)
    suffix = path.suffix.lower()
    ocr = config.OCR_ENABLED if ocr is None else ocr

    if suffix == ".pdf":
        return read_pdf(path, ocr)
    if suffix in IMAGE_TYPES:
        return read_image(path, ocr)
    if suffix in (".xlsx", ".csv"):
        return read_spreadsheet(path)
    if suffix == ".docx":
        return read_docx(path)

    raise ValueError(f"Unsupported file type: {suffix}")
