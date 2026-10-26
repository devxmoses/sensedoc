import docx
import pandas as pd
import pymupdf

# 1. A "photo" of the invoice: render the PDF page to a PNG image
with pymupdf.open("samples/invoice_001.pdf") as pdf:
    pixmap = pdf[0].get_pixmap(dpi=150)
    pixmap.save("samples/invoice_001_photo.png")

# 2. A "scanned" PDF: a PDF that contains only that image, no text
scan = pymupdf.open()
page = scan.new_page()
page.insert_image(page.rect, filename="samples/invoice_001_photo.png")
scan.save("samples/invoice_001_scanned.pdf")

# 3. An Excel price list
prices = pd.DataFrame(
    {
        "Product": ["Coca-Cola Original 50cl PET", "Coke Zero 33cl Can", "Fanta Orange 35cl Can x12"],
        "Price (NGN)": [500, 450, 4200],
    }
)
prices.to_excel("samples/price_list_001.xlsx", index=False, sheet_name="October")

# 4. A Word invoice with a table
document = docx.Document()
document.add_heading("INVOICE", level=1)
document.add_paragraph("Invoice No: INV-2026-0200")
document.add_paragraph("Date: 20 October 2026")
table = document.add_table(rows=2, cols=2)
table.rows[0].cells[0].text = "Product"
table.rows[0].cells[1].text = "Amount"
table.rows[1].cells[0].text = "Pepsi 1L PET"
table.rows[1].cells[1].text = "800.00"
document.save("samples/invoice_002.docx")

print("Made 4 sample files in samples/")
